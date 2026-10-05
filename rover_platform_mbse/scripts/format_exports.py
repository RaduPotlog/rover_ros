#!/usr/bin/env python3

# Copyright 2026 Mechatronics Academy
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Apply the SYS-SR workbook style to the spreadsheets in export/.

MATLAB (export_requirements.m) writes plain cells; this gives them the look of
system/ROVER-A1-ROVER_ROS_SYS-SR-V-002.xlsx: Arial 10, white bold header on dark blue,
hair borders, wrapped top-aligned text, frozen header row and an auto filter.
Verdict / result cells are colour-coded.

Run from WSL after export_requirements:
    uv run --no-project --with openpyxl python scripts/format_exports.py
"""
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXPORT = Path(__file__).resolve().parent.parent / "export"
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E79")
HAIR = Side(style="hair")
BORDER = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
# Colours for status-like cells, matched on the cell text
STATUS_FILL = {
    "Implemented": "FFE2EFDA", "Compliant": "FFE2EFDA", "PASS": "FFE2EFDA",
    "Partial": "FFFFF2CC", "Blocked (TBD)": "FFFFF2CC",
    "Deviation": "FFFCE4D6", "Not Met": "FFF8CBAD", "Gap": "FFF8CBAD",
    "FAIL": "FFF8CBAD", "FAIL (known deviation)": "FFFCE4D6",
    "Not SW": "FFEDEDED",
}
WIDE = {"Requirement", "Rationale", "Acceptance Criteria", "Assessment (as-built, with evidence)",
        "Comments", "Question", "Scope", "Message", "Evidence (src/rover_ros path:line)",
        "Model Verification", "Derived SWRs", "Evidence", "Meaning"}


def width_for(header: str, sheet: str) -> float:
    if sheet == "SYS-SR x Package":
        return 13.0 if header != "SYS-SR" else 12.0
    if header in WIDE:
        return 46.0
    if header in {"Verified By (rover_ros tests)", "Verifies", "Model tests", "rover_ros tests",
                  "Implemented by (architecture)", "Model test", "Trace / Depends On",
                  "Derived from (SYS-SR)"}:
        return 30.0
    return max(10.0, min(24.0, len(header) + 4.0))


def style(path: Path) -> None:
    wb = load_workbook(path)
    for ws in wb.worksheets:
        headers = [c.value or "" for c in ws[1]]
        for col, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width_for(str(header), ws.title)
        for row in ws.iter_rows():
            for cell in row:
                cell.border = BORDER
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if cell.row == 1:
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
                    cell.fill = HEADER_FILL
                else:
                    cell.font = Font(name="Arial", size=10)
                    fill = STATUS_FILL.get(str(cell.value)) if cell.value is not None else None
                    if fill:
                        cell.fill = PatternFill("solid", fgColor=fill)
        if ws.title not in {"Document"}:
            ws.freeze_panes = "B2"
            ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    files = sorted(EXPORT.glob("*.xlsx"))
    for f in files:
        style(f)
    print(f"format_exports: styled {len(files)} workbooks in {EXPORT}")


if __name__ == "__main__":
    main()
